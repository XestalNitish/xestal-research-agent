"""
File tools: file_manager, document_reader, pdf_toolkit, archive_manager, text_diff.

Everything is confined to the workspace (AGENT_WORKSPACE) through safe_path()/safe_output_path(); secret files
(.env, keys, ...) are never read, copied, archived or extracted.

* file_manager     - list/stat/mkdir/copy/move/delete. "delete" never destroys data: the item is moved to
                     <workspace>/_trash/<timestamp>_<name>. Symlinks are refused for copy/move/delete.
* document_reader  - text extraction for txt/md/csv/json/docx/pdf/xlsx with size and output caps.
                     docx needs python-docx, pdf needs pypdf, xlsx needs openpyxl.
* pdf_toolkit      - info / merge / split / extract pages with pypdf (pip install pypdf).
* archive_manager  - create/list/extract zip, tar, tar.gz (stdlib). Extraction validates every member first
                     (zip-slip, absolute paths, symlinks/hardlinks/devices, secret names, encrypted entries) and
                     enforces file count, size and compression-ratio limits while streaming.
* text_diff        - difflib based comparison of two texts or two workspace files.
"""

from __future__ import annotations

import fnmatch
import gzip
import io
import json
import os
import re
import shutil
import stat as stat_mod
import tarfile
import tempfile
import time
import zipfile
import zlib
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any, Dict, List, Tuple

from langchain_core.tools import tool

from ._common import (
    ToolError,
    atomic_write_bytes,
    clamp,
    guard,
    is_secret_path,
    ok_json,
    parse_json_arg,
    require,
    safe_output_path,
    safe_path,
    truncate,
    workspace_root,
)

TRASH_DIR = "_trash"
MAX_COPY_FILES = 5000
MAX_COPY_BYTES = 500_000_000
MAX_READ_BYTES = 50_000_000
MAX_ZIP_PARSE = 200_000_000          # max uncompressed size of an OOXML container we are willing to open
MAX_EXTRACT_FILES = 5000
MAX_EXTRACT_BYTES = 200_000_000
MAX_EXTRACT_FILE = 100_000_000
MAX_RATIO = 200                      # uncompressed / compressed above this (and > 10 MB) is treated as a bomb
MAX_ARCHIVE_INPUT_BYTES = 200_000_000
MAX_PDF_INPUTS = 50
MAX_PDF_PAGES = 2000
MAX_DIFF_LINES = 5000                # difflib is quadratic on repetitive input (~4 s worst case at this size)


def _rel(p: Path) -> str:
    try:
        r = p.relative_to(workspace_root()).as_posix()
    except ValueError:
        return str(p)
    return r or "."


def _no_symlink(path: str) -> None:
    """Reject a path whose last component is a symlink (safe_path would silently follow it)."""
    raw = Path(str(path)).expanduser()
    if not raw.is_absolute():
        raw = workspace_root() / raw
    if raw.is_symlink():
        raise ToolError(f"'{path}' is a symlink; symlinks are not supported for this operation")


def _in_trash(p: Path) -> bool:
    try:
        return p.relative_to(workspace_root()).parts[:1] == (TRASH_DIR,)
    except ValueError:
        return False


def _human(n: int) -> str:
    size = float(n)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024
    return f"{n} B"


def _iso(ts: float) -> str:
    return datetime.fromtimestamp(ts).isoformat(timespec="seconds")


def _split_list(value: str, field: str, limit: int) -> List[str]:
    """Accept a JSON list or a comma/newline separated string."""
    if value is None or not str(value).strip():
        return []
    v = str(value).strip()
    if v.startswith("["):
        items = parse_json_arg(v, expect=(list,))
    else:
        items = [x for x in re.split(r"[,\n]", v)]
    out = [str(i).strip() for i in items if str(i).strip()]
    if len(out) > limit:
        raise ToolError(f"{field}: at most {limit} entries")
    return out


# ---------------------------------------------------------------------------
# file_manager
# ---------------------------------------------------------------------------
def _walk_files(src: Path) -> Tuple[List[Path], List[Path]]:
    """(regular files, folders) below src; symlinks and secret files are skipped; capped."""
    files: List[Path] = []
    dirs: List[Path] = []
    total = 0
    for dirpath, dirnames, filenames in os.walk(src, followlinks=False):
        dirnames[:] = [d for d in dirnames if not (Path(dirpath, d).is_symlink() or is_secret_path(Path(dirpath, d)))]
        dirs += [Path(dirpath, d) for d in dirnames]
        for fn in filenames:
            fp = Path(dirpath, fn)
            if fp.is_symlink() or is_secret_path(fp):
                continue
            total += fp.stat().st_size
            files.append(fp)
            if len(files) > MAX_COPY_FILES:
                raise ToolError(f"more than {MAX_COPY_FILES} files; copy a smaller folder")
            if total > MAX_COPY_BYTES:
                raise ToolError(f"folder larger than {_human(MAX_COPY_BYTES)}")
    return files, dirs


def _trash(p: Path) -> Path:
    root = workspace_root()
    tdir = root / TRASH_DIR
    tdir.mkdir(exist_ok=True)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    target = tdir / f"{stamp}_{p.name}"
    n = 1
    while target.exists():
        target = tdir / f"{stamp}_{n}_{p.name}"
        n += 1
    shutil.move(str(p), str(target))
    return target


@tool
@guard
def file_manager(action: str, path: str = ".", destination: str = "", pattern: str = "*", limit: int = 200,
                 overwrite: bool = False) -> str:
    """Manage files inside the workspace: list, stat, mkdir, copy, move, delete. "delete" is NOT permanent:
    the item is moved to the _trash folder in the workspace (restore it with move). Secret files (.env, keys)
    and paths outside the workspace are refused.

    Args:
        action: list | stat | mkdir | copy | move | delete.
        path: file or folder (list default: the workspace root).
        destination: target path for copy / move (required for those).
        pattern: glob filter for list, e.g. "*.pdf" (default "*").
        limit: max entries returned by list (default 200, max 1000).
        overwrite: for copy/move, replace an existing destination (the old one goes to _trash). Default false.
    """
    act = (action or "").strip().lower()
    if act not in ("list", "stat", "mkdir", "copy", "move", "delete"):
        raise ToolError("action must be list, stat, mkdir, copy, move or delete")
    root = workspace_root()

    if act == "list":
        p = safe_path(path or ".", must_exist=True)
        if not p.is_dir():
            raise ToolError("not a folder; use action=stat for a file")
        n = clamp(limit, 1, 1000)
        entries, hidden, total = [], 0, 0
        for child in sorted(p.iterdir(), key=lambda c: (not c.is_dir(), c.name.lower())):
            if is_secret_path(child):
                hidden += 1
                continue
            if not fnmatch.fnmatch(child.name.lower(), (pattern or "*").lower()):
                continue
            total += 1
            if len(entries) < n:
                try:
                    st = child.lstat()
                except OSError:
                    continue
                kind = "symlink" if child.is_symlink() else "dir" if child.is_dir() else "file"
                entries.append({"name": child.name, "type": kind, "size": None if kind == "dir" else st.st_size,
                                "modified": _iso(st.st_mtime)})
        res: Dict[str, Any] = {"path": _rel(p), "total": total, "shown": len(entries), "entries": entries}
        if hidden:
            res["hidden_secret_files"] = hidden
        return ok_json(res)

    if act == "stat":
        p = safe_path(path, must_exist=True)
        st = p.stat()
        info: Dict[str, Any] = {"path": _rel(p), "type": "dir" if p.is_dir() else "file", "size": st.st_size,
                                "size_human": _human(st.st_size), "modified": _iso(st.st_mtime),
                                "is_symlink": (root / path).is_symlink()}  # absolute `path` replaces root
        if p.is_dir():
            count, size = 0, 0
            for dp, dn, fns in os.walk(p, followlinks=False):
                for fn in fns:
                    count += 1
                    if count > 10_000:
                        break
                    try:
                        size += os.lstat(os.path.join(dp, fn)).st_size
                    except OSError:
                        pass
                if count > 10_000:
                    break
            info.update({"files": count if count <= 10_000 else ">10000", "total_size": size,
                         "total_size_human": _human(size)})
            info.pop("size"), info.pop("size_human")
        else:
            info["extension"] = p.suffix.lower()
        return ok_json(info)

    if act == "mkdir":
        p = safe_path(path)
        if p.exists() and not p.is_dir():
            raise ToolError("a file with that name already exists")
        existed = p.exists()
        p.mkdir(parents=True, exist_ok=True)
        return f"{'Folder already exists' if existed else 'Folder created'}: {_rel(p)}"

    # copy / move / delete
    if act == "delete":
        _no_symlink(path)
        p = safe_path(path, must_exist=True)
        if p == root:
            raise ToolError("refusing to delete the workspace root")
        if p.name == TRASH_DIR and p.parent == root:
            raise ToolError("refusing to delete the trash folder itself")
        if _in_trash(p):
            raise ToolError("item is already in _trash; permanent deletion is not offered by this tool")
        target = _trash(p)
        return f"Moved to trash: {_rel(p)} -> {_rel(target)} (nothing was permanently deleted)"

    if not destination or not destination.strip():
        raise ToolError("destination is required for copy/move")
    _no_symlink(path)
    src = safe_path(path, must_exist=True)
    if src == root:
        raise ToolError("cannot copy/move the workspace root")
    if act == "move" and src.name == TRASH_DIR and src.parent == root:
        raise ToolError("cannot move the trash folder")
    dst = safe_path(destination)
    if dst.exists() and dst.is_dir() and not src.is_dir():
        dst = dst / src.name  # copy file "into" an existing folder
        dst = safe_path(str(dst))
    if dst == src:
        raise ToolError("source and destination are the same")
    if src.is_dir() and (dst == src or src in dst.parents):
        raise ToolError("destination is inside the source folder")
    if dst in src.parents:
        raise ToolError("destination is a parent folder of the source")
    # validate everything that can fail BEFORE the old destination is moved to _trash
    plan: Tuple[List[Path], List[Path]] = ([], [])
    if act == "copy":
        if src.is_file():
            if src.stat().st_size > MAX_COPY_BYTES:
                raise ToolError(f"file larger than {_human(MAX_COPY_BYTES)}")
        else:
            plan = _walk_files(src)
    if dst.exists():
        if not overwrite:
            raise ToolError(f"destination exists: {_rel(dst)} (pass overwrite=true to replace it; the old item goes to _trash)")
        if dst == root or dst in root.parents or dst.name == TRASH_DIR:
            raise ToolError("refusing to overwrite that destination")
        _trash(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)

    if act == "move":
        shutil.move(str(src), str(dst))
        return f"Moved: {_rel(src)} -> {_rel(dst)}"

    # copy
    if src.is_file():
        shutil.copy2(src, dst)
        return f"Copied file: {_rel(src)} -> {_rel(dst)}"
    files, folders = plan
    dst.mkdir(parents=True, exist_ok=True)
    for d in folders:  # keep empty folders
        (dst / d.relative_to(src)).mkdir(parents=True, exist_ok=True)
    for f in files:
        out = dst / f.relative_to(src)
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(f, out)
    return f"Copied folder: {_rel(src)} -> {_rel(dst)} ({len(files)} files; symlinks and secret files skipped)"


# ---------------------------------------------------------------------------
# document_reader
# ---------------------------------------------------------------------------
_TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".xml", ".html", ".htm", ".yaml",
             ".yml", ".ini", ".cfg", ".toml", ".rst", ".py", ".js", ".ts", ".sql", ".tex"}


def _decode(raw: bytes, truncated: bool = False) -> str:
    """UTF-8 (with BOM) / UTF-16 / cp1252 fallback. `truncated`: raw is only a prefix of the file, so a cut
    multi-byte character at its end is not a decoding problem."""
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        if truncated and e.start >= len(raw) - 3:
            return raw[: e.start].decode("utf-8-sig")
    return raw.decode("cp1252", errors="replace")


def _check_ooxml(p: Path) -> None:
    """Refuse zip bombs / non-zip content before handing an OOXML file to a parser."""
    try:
        with zipfile.ZipFile(p) as z:
            infos = z.infolist()
            total = sum(i.file_size for i in infos)
    except zipfile.BadZipFile:
        raise ToolError("file is not a valid Office document (corrupt or wrong extension)")
    if len(infos) > 10_000 or total > MAX_ZIP_PARSE:
        raise ToolError("document container is too large/complex to read safely")


def _read_docx(p: Path, cap: int) -> str:
    docx = require("docx", "python-docx")
    _check_ooxml(p)
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(p))
    out: List[str] = []
    size = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text
            if text.strip():
                out.append(text)
        elif tag == "tbl":
            t = Table(child, d)
            for row in t.rows:
                out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
        size += len(out[-1]) + 1 if out else 0
        if size > cap * 2:
            break
    return "\n".join(out)


def _read_xlsx(p: Path, cap: int) -> str:
    openpyxl = require("openpyxl", "openpyxl")
    _check_ooxml(p)
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out: List[str] = []
    try:
        for ws in wb.worksheets[:20]:
            out.append(f"--- sheet: {ws.title} ---")
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i > 500:
                    out.append("... (more than 500 rows, rest not shown)")
                    break
                out.append(" | ".join("" if v is None else str(v)[:1000] for v in row[:50]))
            if sum(len(x) for x in out) > cap * 2:
                break
        if len(wb.worksheets) > 20:
            out.append(f"... ({len(wb.worksheets) - 20} more sheets not shown)")
    finally:
        wb.close()
    return "\n".join(out)


def _parse_pages(spec: str, total: int, field: str = "pages") -> List[int]:
    """'1-3,5,8-' -> zero-based page indexes (1-based input), order preserved, de-duplicated."""
    result: List[int] = []
    seen = set()
    for part in str(spec).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)?\s*-\s*(\d+)?", part)
        if m and (m.group(1) or m.group(2)):
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else total
        elif re.fullmatch(r"\d+", part):
            a = b = int(part)
        else:
            raise ToolError(f"{field}: cannot parse '{part}' (use e.g. 1-3,5,8-)")
        if a < 1 or b < a or b > total:
            raise ToolError(f"{field}: '{part}' is outside 1..{total}")
        for i in range(a - 1, b):
            if i not in seen:
                seen.add(i)
                result.append(i)
    if not result:
        raise ToolError(f"{field} is empty")
    return result


def _open_pdf(p: Path):
    pypdf = require("pypdf", "pypdf")
    if p.stat().st_size > MAX_READ_BYTES:
        raise ToolError(f"PDF is {_human(p.stat().st_size)}, larger than the {_human(MAX_READ_BYTES)} limit")
    try:
        reader = pypdf.PdfReader(str(p))
    except Exception as e:  # noqa: BLE001 - pypdf raises many types for corrupt files
        raise ToolError(f"cannot open PDF ({type(e).__name__}): {str(e)[:150]}")
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:  # noqa: BLE001
            ok = 0
        if not ok:
            raise ToolError("PDF is password protected")
    return reader


def _read_pdf(p: Path, cap: int, pages: str) -> Tuple[str, str]:
    reader = _open_pdf(p)
    total = len(reader.pages)
    idx = _parse_pages(pages, total) if pages.strip() else list(range(total))
    out: List[str] = []
    size = 0
    shown = 0
    for i in idx:
        try:
            text = reader.pages[i].extract_text() or ""
        except Exception as e:  # noqa: BLE001
            text = f"[could not extract text: {type(e).__name__}]"
        out.append(f"--- page {i + 1} ---\n{text.strip()}")
        shown += 1
        size += len(out[-1])
        if size > cap * 2:
            break
    note = f"pages shown: {shown}/{len(idx)} of {total}"
    body = "\n".join(out)
    if not body.replace("\n", "").strip() or not re.search(r"\w", re.sub(r"--- page \d+ ---", "", body)):
        note += "; NO TEXT LAYER FOUND (probably a scanned PDF, OCR is needed)"
    return body, note


@tool
@guard
def document_reader(path: str, max_chars: int = 20000, pages: str = "") -> str:
    """Extract text from a workspace file: txt, md, csv, tsv, json, log, xml, html, docx, pdf, xlsx (plus common
    source/config text files). Output is capped. docx needs python-docx, pdf needs pypdf, xlsx needs openpyxl.
    Scanned PDFs have no text layer and are reported as such (no OCR).

    Args:
        path: file inside the workspace.
        max_chars: maximum characters returned (default 20000, min 100, max 100000).
        pages: PDF only: page selection such as "1-3,7" (default all pages, until max_chars is reached).
    """
    p = safe_path(path, must_exist=True)
    if not p.is_file():
        raise ToolError("not a file")
    size = p.stat().st_size
    cap = clamp(max_chars, 100, 100_000)
    ext = p.suffix.lower()
    note = ""
    if ext in (".docx", ".xlsx", ".pdf") and size > MAX_READ_BYTES:
        raise ToolError(f"file is {_human(size)}, larger than the {_human(MAX_READ_BYTES)} limit")
    if ext == ".docx":
        text = _read_docx(p, cap)
    elif ext == ".xlsx":
        text = _read_xlsx(p, cap)
    elif ext == ".pdf":
        text, note = _read_pdf(p, cap, pages)
    elif ext in _TEXT_EXT:
        with p.open("rb") as f:
            raw = f.read(cap * 4 + 8)
        text = _decode(raw, truncated=size > len(raw))
        if size > len(raw):
            note = f"file is {_human(size)}, only the beginning was read"
        elif ext == ".json":
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except ValueError as e:
                note = f"not valid JSON ({e}); raw text shown"
    elif ext in (".doc", ".xls", ".ppt", ".pptx", ".odt"):
        raise ToolError(f"{ext} files are not supported (convert to docx/pdf/xlsx/txt first)")
    else:
        raise ToolError(f"unsupported file type '{ext or '(none)'}'. Supported: txt md csv tsv json docx pdf xlsx "
                        "and common text/config/source files")
    total_chars = len(text)
    shown = truncate(text, cap, marker="\n...[truncated, use max_chars or pages to read more]")
    header = f"[{p.name} | {ext.lstrip('.')} | {_human(size)} | {total_chars} chars extracted"
    header += f" | {note}]" if note else "]"
    if not text.strip():
        return header + "\n(no text content)"
    return header + "\n" + shown


# ---------------------------------------------------------------------------
# pdf_toolkit
# ---------------------------------------------------------------------------
def _pdf_out(output_path: str, default: str) -> Path:
    out = safe_output_path(output_path or default, default)
    if out.suffix.lower() != ".pdf":
        raise ToolError("output_path must end with .pdf")
    return out


def _write_pdf(writer, out: Path) -> int:
    buf = io.BytesIO()
    writer.write(buf)
    if out.exists():  # never silently destroy a different existing file: it goes to _trash
        _trash(out)
    atomic_write_bytes(out, buf.getvalue())
    return len(buf.getvalue())


@tool
@guard
def pdf_toolkit(action: str, path: str, output_path: str = "", pages: str = "") -> str:
    """PDF operations with pypdf (pip install pypdf): info, merge, split, extract. Password protected PDFs are
    refused. Inputs and outputs stay in the workspace; an existing output file is moved to _trash
    before it is replaced; inputs are never modified.

    Args:
        action: info | merge | split | extract.
        path: the PDF. For merge: several PDFs as a comma separated string or JSON list, merged in that order.
        output_path: info: unused. merge/extract: output .pdf (default merged.pdf / extract.pdf). split: output
            FOLDER for the parts (default <name>_split).
        pages: extract: pages to keep, e.g. "1-3,7,9-" (1-based, required). split: ranges that each become one
            file, e.g. "1-3,4-6" (default: one file per page, max 200 pages).
    """
    act = (action or "").strip().lower()
    if act not in ("info", "merge", "split", "extract"):
        raise ToolError("action must be info, merge, split or extract")
    pypdf = require("pypdf", "pypdf")

    if act == "merge":
        names = _split_list(path, "path", MAX_PDF_INPUTS)
        if len(names) < 2:
            raise ToolError("merge needs at least two PDFs in path (comma separated or JSON list)")
        out = _pdf_out(output_path, "merged.pdf")
        writer = pypdf.PdfWriter()
        total = 0
        for n in names:
            src = safe_path(n, must_exist=True)
            if src == out:
                raise ToolError("output_path must differ from the input files")
            if src.suffix.lower() != ".pdf":
                raise ToolError(f"not a PDF: {n}")
            reader = _open_pdf(src)
            total += len(reader.pages)
            if total > MAX_PDF_PAGES:
                raise ToolError(f"more than {MAX_PDF_PAGES} pages in total")
            for page in reader.pages:
                writer.add_page(page)
        size = _write_pdf(writer, out)
        return f"Merged {len(names)} PDFs ({total} pages) -> {_rel(out)} ({_human(size)})"

    src = safe_path(path, must_exist=True)
    if src.suffix.lower() != ".pdf" or not src.is_file():
        raise ToolError("path must be a .pdf file")
    reader = _open_pdf(src)
    total = len(reader.pages)

    if act == "info":
        meta = {}
        try:
            for k, v in (reader.metadata or {}).items():
                if v is not None:
                    meta[str(k).lstrip("/")] = str(v)[:200]
        except Exception:  # noqa: BLE001 - broken metadata must not break info
            meta = {}
        first = reader.pages[0].mediabox if total else None
        sample = ""
        if total:
            try:
                sample = (reader.pages[0].extract_text() or "")[:200]
            except Exception:  # noqa: BLE001
                sample = ""
        return ok_json({"path": _rel(src), "pages": total, "size": _human(src.stat().st_size),
                        "page_size_points": [round(float(first.width)), round(float(first.height))] if first else None,
                        "metadata": meta, "has_text_layer": bool(sample.strip()),
                        "first_page_preview": sample.strip()})

    if total > MAX_PDF_PAGES:
        raise ToolError(f"PDF has more than {MAX_PDF_PAGES} pages")

    if act == "extract":
        if not pages.strip():
            raise ToolError("pages is required for extract, e.g. '1-3,7'")
        idx = _parse_pages(pages, total)
        out = _pdf_out(output_path, "extract.pdf")
        if out == src:
            raise ToolError("output_path must differ from the input file")
        writer = pypdf.PdfWriter()
        for i in idx:
            writer.add_page(reader.pages[i])
        size = _write_pdf(writer, out)
        return f"Extracted {len(idx)} page(s) -> {_rel(out)} ({_human(size)})"

    # split
    if pages.strip():
        groups = [_parse_pages(part, total, "pages") for part in pages.split(",") if part.strip()]
    else:
        if total > 200:
            raise ToolError("PDF has more than 200 pages: pass pages='1-50,51-100,...' to split into chunks")
        groups = [[i] for i in range(total)]
    if len(groups) > 200:
        raise ToolError("more than 200 parts requested")
    folder = safe_path(output_path or f"{src.stem}_split")
    if folder.exists() and not folder.is_dir():
        raise ToolError("output_path must be a folder")
    folder.mkdir(parents=True, exist_ok=True)
    created = []
    for n, g in enumerate(groups, start=1):
        label = f"p{g[0] + 1}" if len(g) == 1 else f"p{g[0] + 1}-{g[-1] + 1}"
        target = folder / f"{src.stem}_part{n:03d}_{label}.pdf"
        writer = pypdf.PdfWriter()
        for i in g:
            writer.add_page(reader.pages[i])
        _write_pdf(writer, target)
        created.append(_rel(target))
    return ok_json({"split_into": len(created), "folder": _rel(folder), "files": created})


# ---------------------------------------------------------------------------
# archive_manager
# ---------------------------------------------------------------------------
_BAD_NAME_CHARS = re.compile(r'[<>"|?*\x00-\x1f:]')
_RESERVED_WIN = {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}


def _archive_kind(p: Path) -> str:
    n = p.name.lower()
    if n.endswith(".zip"):
        return "zip"
    if n.endswith((".tar.gz", ".tgz")):
        return "tar.gz"
    if n.endswith(".tar"):
        return "tar"
    raise ToolError("archive must end with .zip, .tar, .tar.gz or .tgz")


def _clean_member_name(name: str) -> Tuple[str, ...]:
    """Validate an archive member name; returns its path parts. Raises ToolError on anything suspicious."""
    n = name.replace("\\", "/")
    if not n or n.startswith("/") or re.match(r"^[A-Za-z]:", n):
        raise ToolError(f"unsafe absolute path in archive: {name!r}")
    parts = [x for x in PurePosixPath(n).parts if x not in (".", "")]
    if not parts:
        raise ToolError(f"empty member name: {name!r}")
    if ".." in parts:
        raise ToolError(f"path traversal (zip-slip) in archive member: {name!r}")
    for part in parts:
        if _BAD_NAME_CHARS.search(part):
            raise ToolError(f"illegal characters in archive member name: {name!r}")
        if part.split(".")[0].lower() in _RESERVED_WIN or part.endswith((" ", ".")):
            raise ToolError(f"reserved/unsafe file name in archive: {name!r}")
    if is_secret_path(Path(*parts)):
        raise ToolError(f"archive contains a secret/credential file name, refusing to extract: {name!r}")
    return tuple(parts)


def _list_zip(z: zipfile.ZipFile) -> List[Dict[str, Any]]:
    return [{"name": i.filename, "size": i.file_size, "compressed": i.compress_size, "dir": i.is_dir()}
            for i in z.infolist()]


def _tar_member_kind(m: tarfile.TarInfo) -> str:
    if m.isdir():
        return "dir"
    if m.isreg():
        return "file"
    return "other"


def _check_limits(count: int, total: int, largest: int, archive_size: int) -> None:
    if count > MAX_EXTRACT_FILES:
        raise ToolError(f"archive has {count} entries (limit {MAX_EXTRACT_FILES})")
    if total > MAX_EXTRACT_BYTES:
        raise ToolError(f"archive would extract to {_human(total)} (limit {_human(MAX_EXTRACT_BYTES)})")
    if largest > MAX_EXTRACT_FILE:
        raise ToolError(f"a single file is {_human(largest)} (limit {_human(MAX_EXTRACT_FILE)})")
    if total > 10_000_000 and total / max(archive_size, 1) > MAX_RATIO:
        raise ToolError(f"suspicious compression ratio (>{MAX_RATIO}:1): possible zip bomb")


def _copy_limited(src, dst_path: Path, budget: List[int], declared: int) -> None:
    """Stream src -> dst_path, enforcing the real number of bytes (declared sizes in archives can lie)."""
    written = 0
    with open(dst_path, "wb") as out:
        while True:
            chunk = src.read(65536)
            if not chunk:
                break
            written += len(chunk)
            budget[0] -= len(chunk)
            if written > declared + 1 or budget[0] < 0 or written > MAX_EXTRACT_FILE:
                raise ToolError("archive entry is larger than declared or exceeds the size limit (zip bomb?)")
            out.write(chunk)


def _extract(arc: Path, kind: str, dest: Path, overwrite: bool) -> Tuple[int, int]:
    asize = arc.stat().st_size
    entries: List[Tuple[Tuple[str, ...], str, int, Any]] = []  # (parts, "file"/"dir", size, handle)
    opener = None
    try:
        if kind == "zip":
            z = zipfile.ZipFile(arc)
            opener = z
            for info in z.infolist():
                mode = (info.external_attr >> 16) & 0xFFFF
                ftype = stat_mod.S_IFMT(mode)  # 0 when the zip only stores permission bits (e.g. writestr)
                if info.flag_bits & 0x1:
                    raise ToolError("archive contains encrypted entries (not supported)")
                if ftype and ftype not in (stat_mod.S_IFREG, stat_mod.S_IFDIR):
                    raise ToolError(f"archive contains a symlink or special file: {info.filename!r}")
                parts = _clean_member_name(info.filename)
                entries.append((parts, "dir" if info.is_dir() else "file", info.file_size, info))
        else:
            t = tarfile.open(arc, "r:*")
            opener = t
            declared = 0
            for m in t:  # iterating a .tar.gz decompresses it, so stop as soon as a limit is exceeded
                k = _tar_member_kind(m)
                if k == "other":
                    raise ToolError(f"archive contains a symlink, hardlink or special file: {m.name!r}")
                entries.append((_clean_member_name(m.name), k, m.size if k == "file" else 0, m))
                declared += m.size if k == "file" else 0
                if len(entries) > MAX_EXTRACT_FILES or declared > MAX_EXTRACT_BYTES or m.size > MAX_EXTRACT_FILE:
                    break
        files = [e for e in entries if e[1] == "file"]
        _check_limits(len(entries), sum(e[2] for e in files), max((e[2] for e in files), default=0), asize)

        dest.parent.mkdir(parents=True, exist_ok=True)
        tmp = Path(tempfile.mkdtemp(prefix=f".{dest.name}.extracting-", dir=str(dest.parent)))
        budget = [MAX_EXTRACT_BYTES]
        try:
            tmp_root = tmp.resolve()
            for parts, k, size, handle in entries:
                target = (tmp_root.joinpath(*parts))
                resolved = target.resolve()
                if resolved != tmp_root and tmp_root not in resolved.parents:  # defence in depth
                    raise ToolError(f"path traversal in archive member: {'/'.join(parts)!r}")
                if k == "dir":
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.exists():
                    raise ToolError(f"duplicate entry in archive: {'/'.join(parts)!r}")
                if kind == "zip":
                    with opener.open(handle) as src:  # type: ignore[union-attr]
                        _copy_limited(src, target, budget, size)
                else:
                    src = opener.extractfile(handle)  # type: ignore[union-attr]
                    if src is None:
                        raise ToolError(f"cannot read member {'/'.join(parts)!r}")
                    with src:
                        _copy_limited(src, target, budget, size)
            if dest.exists():
                if any(dest.iterdir()) and not overwrite:
                    raise ToolError(f"destination folder is not empty: {_rel(dest)} (pass overwrite=true; the old "
                                    "folder will be moved to _trash)")
                if any(dest.iterdir()):
                    _trash(dest)
                else:
                    dest.rmdir()
            os.chmod(tmp, 0o755)  # mkdtemp creates 0700
            os.replace(tmp, dest)
        except BaseException:
            shutil.rmtree(tmp, ignore_errors=True)
            raise
        return len(files), sum(e[2] for e in files)
    except (zipfile.BadZipFile, zipfile.LargeZipFile, tarfile.TarError, EOFError, zlib.error, gzip.BadGzipFile,
            NotImplementedError) as e:
        raise ToolError(f"cannot read archive ({type(e).__name__}): {str(e)[:150]}")
    finally:
        if opener is not None:
            opener.close()


def _gather_sources(sources: List[str], out: Path) -> Tuple[List[Tuple[Path, str]], int]:
    """(file, arcname) pairs; symlinks / secret files / the output archive itself are skipped. Returns also #skipped."""
    pairs: List[Tuple[Path, str]] = []
    skipped = 0
    total = 0
    seen = set()
    for s in sources:
        _no_symlink(s)
        src = safe_path(s, must_exist=True)
        if src == workspace_root():
            raise ToolError("refusing to archive the whole workspace root; list sub-folders instead")
        base = src.parent
        walk = [src] if src.is_file() else None
        if walk is None:
            walk = []
            for dp, dn, fns in os.walk(src, followlinks=False):
                keep = []
                for d in dn:
                    dpth = Path(dp, d)
                    if dpth.is_symlink() or is_secret_path(dpth):
                        skipped += 1
                    else:
                        keep.append(d)
                dn[:] = keep
                walk += [Path(dp, f) for f in fns]
                if not fns and not dn:
                    walk.append(Path(dp))  # keep empty folders
        for f in walk:
            if f.is_symlink() or is_secret_path(f) or f == out:
                skipped += 1
                continue
            arc = f.relative_to(base).as_posix()
            if arc in seen:
                continue
            seen.add(arc)
            if f.is_file():
                total += f.stat().st_size
            pairs.append((f, arc))
            if len(pairs) > MAX_EXTRACT_FILES:
                raise ToolError(f"more than {MAX_EXTRACT_FILES} files to archive")
            if total > MAX_ARCHIVE_INPUT_BYTES:
                raise ToolError(f"input larger than {_human(MAX_ARCHIVE_INPUT_BYTES)}")
    if not pairs:
        raise ToolError("nothing to archive (empty folder or only secret/symlink files)")
    return pairs, skipped


@tool
@guard
def archive_manager(action: str, path: str, output_path: str = "", overwrite: bool = False) -> str:
    """Create, list or safely extract zip / tar / tar.gz archives inside the workspace.
    Extraction first validates EVERY entry and refuses the whole archive on path traversal (zip-slip), absolute
    paths, symlinks/hardlinks/devices, secret file names (.env, keys), encrypted entries, too many files
    (5000), too much data (200 MB) or an absurd compression ratio (zip bomb); nothing is written in that case.
    Archives never include secret files or symlinks.

    Args:
        action: create | list | extract.
        path: create: file(s)/folder(s) to pack, comma separated or JSON list. list/extract: the archive file.
        output_path: create: archive to write (.zip, .tar, .tar.gz or .tgz; default archive.zip). extract:
            destination folder (default: archive name without extension).
        overwrite: replace an existing archive (create) or non-empty folder (extract; the old one goes to _trash).
    """
    act = (action or "").strip().lower()
    if act not in ("create", "list", "extract"):
        raise ToolError("action must be create, list or extract")

    if act == "create":
        sources = _split_list(path, "path", 100)
        if not sources:
            raise ToolError("path is empty: give the files/folders to archive")
        out = safe_output_path(output_path or "archive.zip", "archive.zip")
        kind = _archive_kind(out)
        if out.exists() and not overwrite:
            raise ToolError(f"{_rel(out)} already exists (pass overwrite=true to replace)")
        pairs, skipped = _gather_sources(sources, out)
        fd, tmp_name = tempfile.mkstemp(dir=str(out.parent), prefix=out.name + ".", suffix=".tmp")
        os.close(fd)
        tmp = Path(tmp_name)
        try:
            if kind == "zip":
                with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, strict_timestamps=False) as z:
                    for f, arc in pairs:
                        if f.is_dir():
                            z.writestr(arc.rstrip("/") + "/", "")
                        else:
                            z.write(f, arc)
            else:
                with tarfile.open(tmp, "w:gz" if kind == "tar.gz" else "w") as t:
                    for f, arc in pairs:
                        t.add(f, arcname=arc, recursive=False, filter=_tar_filter)
            if out.exists():
                _trash(out)
            os.replace(tmp, out)
        except BaseException:
            tmp.unlink(missing_ok=True)
            raise
        n_files = sum(1 for f, _ in pairs if f.is_file())
        msg = f"Created {kind} archive {_rel(out)}: {n_files} files, {_human(out.stat().st_size)}"
        return msg + (f" ({skipped} secret/symlink/self entries skipped)" if skipped else "")

    arc = safe_path(path, must_exist=True)
    if not arc.is_file():
        raise ToolError("path must be an archive file")
    kind = _archive_kind(arc)

    if act == "list":
        try:
            if kind == "zip":
                with zipfile.ZipFile(arc) as z:
                    items = _list_zip(z)
            else:
                with tarfile.open(arc, "r:*") as t:
                    items, declared = [], 0
                    for m in t:  # a .tar.gz is decompressed while iterating: stop early on huge archives
                        items.append({"name": m.name, "size": m.size, "dir": m.isdir(),
                                      "type": "file" if m.isreg() else "dir" if m.isdir() else "SPECIAL"})
                        declared += m.size
                        if len(items) > 20_000 or declared > MAX_EXTRACT_BYTES:
                            break
        except (zipfile.BadZipFile, tarfile.TarError, EOFError, zlib.error, gzip.BadGzipFile) as e:
            raise ToolError(f"cannot read archive ({type(e).__name__}): {str(e)[:150]}")
        total = sum(i["size"] for i in items)
        return ok_json({"archive": _rel(arc), "entries": len(items), "total_uncompressed": _human(total),
                        "shown": items[:200], "truncated": len(items) > 200})

    # extract
    default_dest = re.sub(r"(\.tar\.gz|\.tgz|\.tar|\.zip)$", "", arc.name, flags=re.I) or "extracted"
    dest = safe_path(output_path or str(arc.parent / default_dest))
    if dest == workspace_root() or dest in workspace_root().parents:
        raise ToolError("choose a sub-folder as destination, not the workspace root")
    if dest.exists() and not dest.is_dir():
        raise ToolError("destination exists and is not a folder")
    n_files, n_bytes = _extract(arc, kind, dest, overwrite)
    return f"Extracted {n_files} files ({_human(n_bytes)}) to {_rel(dest)}"


def _tar_filter(ti: tarfile.TarInfo) -> tarfile.TarInfo:
    ti.uid = ti.gid = 0
    ti.uname = ti.gname = ""
    ti.mode = 0o755 if ti.isdir() else 0o644
    return ti


# ---------------------------------------------------------------------------
# text_diff
# ---------------------------------------------------------------------------
def _diff_input(text: str, file_path: str, label: str) -> Tuple[List[str], str]:
    if file_path and file_path.strip():
        p = safe_path(file_path, must_exist=True)
        if not p.is_file():
            raise ToolError(f"{label}: not a file")
        if p.stat().st_size > 5_000_000:
            raise ToolError(f"{label}: file larger than 5 MB")
        content = _decode(p.read_bytes())
        name = _rel(p)
    else:
        content, name = text or "", label
    lines = content.splitlines()
    if len(lines) > MAX_DIFF_LINES:
        raise ToolError(f"{label}: more than {MAX_DIFF_LINES} lines")
    return lines, name


@tool
@guard
def text_diff(text_a: str = "", text_b: str = "", path_a: str = "", path_b: str = "", mode: str = "unified",
              context: int = 3, ignore_whitespace: bool = False) -> str:
    """Compare two texts or two workspace files line by line (difflib) and return a unified diff or a summary.
    For each side give either the text or a file path (the path wins when both are given).

    Args:
        text_a: first text (original).
        text_b: second text (changed).
        path_a: workspace file to use instead of text_a.
        path_b: workspace file to use instead of text_b.
        mode: unified (default, full diff) | summary (counts of added/removed/changed lines only).
        context: unified diff context lines (0-20, default 3).
        ignore_whitespace: compare with leading/trailing whitespace removed and runs of spaces collapsed
            (the diff then shows the normalised lines).
    Line endings (CRLF vs LF) and a missing final newline are not treated as differences. Max 5000 lines per side.
    """
    import difflib

    m = (mode or "unified").strip().lower()
    if m not in ("unified", "summary"):
        raise ToolError("mode must be unified or summary")
    a, name_a = _diff_input(text_a, path_a, "text_a")
    b, name_b = _diff_input(text_b, path_b, "text_b")
    if not (a or b) and not (path_a or path_b):
        raise ToolError("nothing to compare: give text_a/text_b or path_a/path_b")
    if ignore_whitespace:
        a = [re.sub(r"\s+", " ", x).strip() for x in a]
        b = [re.sub(r"\s+", " ", x).strip() for x in b]
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    added = removed = changed = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "insert":
            added += j2 - j1
        elif tag == "delete":
            removed += i2 - i1
        elif tag == "replace":
            changed += max(i2 - i1, j2 - j1)
            removed += max(0, (i2 - i1) - (j2 - j1))
            added += max(0, (j2 - j1) - (i2 - i1))
    if a == b:
        return f"No differences ({len(a)} lines compared)."
    stats = (f"lines: {len(a)} -> {len(b)} | added {added}, removed {removed}, changed {changed} | "
             f"similarity {sm.ratio():.0%}")
    if m == "summary":
        return stats
    diff = list(difflib.unified_diff(a, b, fromfile=name_a, tofile=name_b, n=clamp(context, 0, 20), lineterm=""))
    return stats + "\n" + truncate("\n".join(diff), 20000, "\n...[diff truncated]")


FILE_OPS_TOOLS = [file_manager, document_reader, pdf_toolkit, archive_manager, text_diff]
