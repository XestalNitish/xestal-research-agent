"""
Document generation tools: docx_creator, pdf_creator, pptx_creator, xlsx_creator, markdown_converter.

Input is a structured text argument. Either JSON, or simple Markdown-like text:
    # Heading / ## Sub heading     - bullet     1. numbered     | a | b | table rows     blank line = new paragraph
JSON for docx/pdf: {"title": "...", "blocks": [{"type": "heading", "text": "..", "level": 1},
    {"type": "paragraph", "text": ".."}, {"type": "bullets", "items": [..]}, {"type": "numbered", "items": [..]},
    {"type": "table", "rows": [[..], [..]]}]}  (a bare list of blocks also works).
Every file is written into the workspace via safe_output_path and the absolute path is returned.

Optional packages: python-docx (docx_creator), reportlab (pdf_creator), python-pptx (pptx_creator),
openpyxl (xlsx_creator), markdown (markdown_converter; a small built-in converter is used when missing).

Inline **bold**, *italic* and `code` are honoured in docx/pdf paragraphs, list items and table cells.

PDF and Devanagari: reportlab's built-in fonts have no Devanagari glyphs. pdf_creator looks for a Unicode font
(argument font_path, env AGENT_PDF_FONT, then Nirmala UI / Mangal on Windows, Noto Sans Devanagari / Lohit on
Linux/macOS) that really contains every needed glyph. If the text needs glyphs no font provides, the tool fails
loudly instead of writing a PDF full of black boxes. Note: reportlab does not do complex-script shaping, so conjuncts may look
imperfect; for perfect Hindi PDFs use markdown_converter -> HTML -> browser "Print to PDF".
"""

from __future__ import annotations

import csv
import glob
import html
import io
import json
import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, List

from langchain_core.tools import tool

from ._common import (
    ToolError,
    atomic_write_bytes,
    atomic_write_text,
    guard,
    parse_json_arg,
    require,
    safe_output_path,
    safe_path,
)

MAX_CONTENT_CHARS = 500_000
MAX_BLOCKS = 5000
MAX_TABLE_CELLS = 100_000
MAX_SLIDES = 200

_DEVANAGARI = re.compile(r"[\u0900-\u097F]")
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")
_INLINE = re.compile(r"(\*\*[^*\n]+\*\*|(?<!\*)\*[^*\n]+\*(?!\*)|`[^`\n]+`)")


def _clean(text: Any) -> str:
    """str() + remove control characters that XML based formats (docx/xlsx/pptx) reject."""
    return _CTRL.sub("", str(text))


def _runs(text: str) -> List[tuple]:
    """Split text into (segment, style) pairs; style in '', 'b', 'i', 'c' (inline markdown)."""
    out: List[tuple] = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], "b"))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], "c"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], "i"))
        else:
            out.append((part, ""))
    return out


# ---------------------------------------------------------------------------
# shared parsing
# ---------------------------------------------------------------------------
def _check_content(content: str) -> str:
    if not isinstance(content, str) or not content.strip():
        raise ToolError("content is empty")
    if len(content) > MAX_CONTENT_CHARS:
        raise ToolError(f"content longer than {MAX_CONTENT_CHARS} characters")
    return content


def _text_to_blocks(text: str) -> List[Dict[str, Any]]:
    blocks: List[Dict[str, Any]] = []
    para: List[str] = []

    def flush():
        if para:
            blocks.append({"type": "paragraph", "text": " ".join(para)})
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        s = line.strip()
        m_h = re.match(r"^(#{1,3})\s+(.*)$", s)
        m_b = re.match(r"^[-*•]\s+(.*)$", s)
        m_n = re.match(r"^\d+[.)]\s+(.*)$", s)
        if not s:
            flush()
        elif m_h:
            flush()
            blocks.append({"type": "heading", "level": len(m_h.group(1)), "text": m_h.group(2)})
        elif m_b or m_n:
            flush()
            kind, item = ("bullets", m_b.group(1)) if m_b else ("numbered", m_n.group(1))
            if blocks and blocks[-1]["type"] == kind:
                blocks[-1]["items"].append(item)
            else:
                blocks.append({"type": kind, "items": [item]})
        elif s.startswith("|") and s.endswith("|"):
            flush()
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if blocks and blocks[-1]["type"] == "table":
                blocks[-1]["rows"].append(cells)
            else:
                blocks.append({"type": "table", "rows": [cells]})
        else:
            para.append(s)
    flush()
    return blocks


def _try_json(s: str) -> Any:
    """Parsed JSON (dict/list) when `s` looks like JSON, else None (plain text). Broken '{...' JSON raises."""
    if s[0] not in "{[":
        return None
    try:
        return parse_json_arg(s)
    except ToolError:
        if s[0] == "{":
            raise
        return None


def _parse_blocks(content: str, title: str = "") -> tuple:
    """Return (title, blocks) from JSON or markdown-like text; validates structure."""
    _check_content(content)
    s = content.strip()
    blocks: List[Dict[str, Any]]
    data = _try_json(s)
    if data is not None:
        if isinstance(data, dict):
            title = title or str(data.get("title", ""))
            blocks = data.get("blocks")  # type: ignore[assignment]
            if not isinstance(blocks, list):
                raise ToolError('JSON object needs a "blocks" list')
        else:
            blocks = data
    else:
        blocks = _text_to_blocks(s)
    title = _clean(title)
    if not blocks:
        raise ToolError("no content blocks found")
    if len(blocks) > MAX_BLOCKS:
        raise ToolError(f"more than {MAX_BLOCKS} blocks")
    clean: List[Dict[str, Any]] = []
    cells = 0
    for b in blocks:
        if not isinstance(b, dict):
            raise ToolError("every block must be an object")
        t = str(b.get("type", "paragraph")).lower()
        if t == "heading":
            level = int(str(b.get("level", 1))) if str(b.get("level", 1)).isdigit() else 1
            clean.append({"type": "heading", "level": min(max(level, 1), 3), "text": _clean(b.get("text", ""))})
        elif t == "paragraph":
            clean.append({"type": "paragraph", "text": _clean(b.get("text", ""))})
        elif t in ("bullets", "numbered"):
            items = b.get("items")
            if not isinstance(items, list):
                raise ToolError(f"{t} block needs an items list")
            clean.append({"type": t, "items": [_clean(i) for i in items]})
        elif t == "table":
            rows = b.get("rows")
            if (not isinstance(rows, list) or not rows or not all(isinstance(r, list) for r in rows)
                    or not any(rows)):
                raise ToolError("table block needs rows: a non-empty list of lists")
            cells += sum(len(r) for r in rows)
            if cells > MAX_TABLE_CELLS:
                raise ToolError("tables too large")
            clean.append({"type": "table", "rows": [[_clean(c) for c in r] for r in rows]})
        else:
            raise ToolError(f"unknown block type '{t}'")
    return title, clean


def _out(path: str, default: str, suffix: str) -> Path:
    target = safe_path(path or default)  # validates the sandbox without creating anything
    if not (target.is_dir() or target.suffix.lower() == suffix):  # a folder means "default name inside it"
        raise ToolError(f"output_path must end with {suffix}")
    p = safe_output_path(path or default, default)
    if p.suffix.lower() != suffix:
        raise ToolError(f"output_path must end with {suffix}")
    return p


# ---------------------------------------------------------------------------
# docx_creator
# ---------------------------------------------------------------------------
def _add_runs(paragraph, text: str, bold: bool = False) -> None:
    for seg, st in _runs(text):
        run = paragraph.add_run(seg)
        run.bold = True if (bold or st == "b") else None
        run.italic = True if st == "i" else None
        if st == "c":
            run.font.name = "Consolas"


@tool
@guard
def docx_creator(content: str, output_path: str = "document.docx", title: str = "") -> str:
    """Create a Word (.docx) document from structured text and save it in the workspace. Needs: pip install python-docx.

    Args:
        content: JSON ({"title":..,"blocks":[{"type":"heading|paragraph|bullets|numbered|table",...}]}) or
            Markdown-like text (# headings, - bullets, 1. numbered, | table | rows |).
        output_path: destination .docx path inside the workspace (default document.docx).
        title: optional document title shown at the top.
    """
    title, blocks = _parse_blocks(content, title)
    out = _out(output_path, "document.docx", ".docx")
    docx = require("docx", "python-docx")
    doc = docx.Document()
    if title:
        doc.add_heading(title, level=0)
    for b in blocks:
        t = b["type"]
        if t == "heading":
            doc.add_heading(b["text"], level=b["level"])
        elif t == "paragraph":
            _add_runs(doc.add_paragraph(), b["text"])
        elif t in ("bullets", "numbered"):
            style = "List Bullet" if t == "bullets" else "List Number"
            for item in b["items"]:
                _add_runs(doc.add_paragraph(style=style), item)
        else:
            rows = b["rows"]
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    _add_runs(table.cell(i, j).paragraphs[0], c, bold=(i == 0))
    buf = io.BytesIO()
    doc.save(buf)
    atomic_write_bytes(out, buf.getvalue())
    return f"DOCX saved: {out} ({len(blocks)} blocks)"


# ---------------------------------------------------------------------------
# pdf_creator
# ---------------------------------------------------------------------------
_FONT_CANDIDATES = [
    r"C:\Windows\Fonts\Nirmala.ttf", r"C:\Windows\Fonts\NirmalaUI.ttf", r"C:\Windows\Fonts\mangal.ttf",
    r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\segoeui.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
    "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]
_FONT_GLOBS = ["/usr/share/fonts/**/*Devanagari*.ttf", "/usr/share/fonts/**/*Sanskrit*.ttf",
               "/usr/share/fonts/**/*Lohit-Devanagari*.ttf", "/usr/share/fonts/**/*Mukta*.ttf",
               "/usr/local/share/fonts/**/*Devanagari*.ttf", os.path.expanduser("~/.fonts/**/*Devanagari*.ttf"),
               os.path.expanduser("~/.local/share/fonts/**/*Devanagari*.ttf")]


def _font_candidates(font_path: str) -> List[str]:
    if font_path:
        return [str(safe_path(font_path, must_exist=True))]
    out: List[str] = []
    env = os.environ.get("AGENT_PDF_FONT", "").strip()
    if env and Path(env).is_file():
        out.append(env)
    out += [c for c in _FONT_CANDIDATES if Path(c).is_file()]
    for g in _FONT_GLOBS:
        out += sorted(glob.glob(g, recursive=True))
    return list(dict.fromkeys(out))


def _missing_glyphs(font_file: str, chars: set) -> set:
    """Characters in `chars` the TTF has no glyph for (checked with fontTools, a reportlab/matplotlib dependency)."""
    try:
        from fontTools.ttLib import TTFont as FTFont  # type: ignore

        cmap = FTFont(font_file, fontNumber=0, lazy=True).getBestCmap() or {}
    except Exception:  # noqa: BLE001 - unreadable font: treat everything as missing
        return set(chars)
    return {ch for ch in chars if ord(ch) not in cmap}


def _needs_unicode_font(text: str) -> set:
    """Characters that the built-in Helvetica (WinAnsi / cp1252) cannot show."""
    need = set()
    for ch in set(text):
        if ch.isspace() or unicodedata.category(ch) in ("Cc", "Cf"):
            continue
        try:
            ch.encode("cp1252")
        except UnicodeEncodeError:
            need.add(ch)
    return need


def _pdf_markup(text: str) -> str:
    """Escape text for reportlab Paragraph and turn inline markdown into <b>/<i>/<font> tags."""
    parts = []
    for seg, st in _runs(text):
        e = html.escape(seg).replace("\n", "<br/>")
        parts.append({"b": f"<b>{e}</b>", "i": f"<i>{e}</i>", "c": f'<font face="Courier">{e}</font>'}.get(st, e))
    return "".join(parts)


@tool
@guard
def pdf_creator(content: str, output_path: str = "document.pdf", title: str = "", font_path: str = "") -> str:
    """Create a PDF from structured text and save it in the workspace. Needs: pip install reportlab.
    Devanagari/Hindi (and other non-Latin) text needs a Unicode TTF font containing those glyphs (Nirmala UI /
    Mangal on Windows, Noto Sans Devanagari on Linux); if none is found the tool fails instead of producing
    black boxes. Pass font_path or set AGENT_PDF_FONT. Limitation: reportlab does not shape complex scripts, so
    some Devanagari vowel signs/conjuncts can look imperfect (the result says so); for exact Hindi output use
    markdown_converter and print the HTML to PDF from a browser.

    Args:
        content: JSON ({"title":..,"blocks":[...]}) or Markdown-like text (# headings, - bullets, 1. numbered, | tables |).
        output_path: destination .pdf path inside the workspace (default document.pdf).
        title: optional title at the top of the first page.
        font_path: optional .ttf path (inside the workspace) used for all text.
    """
    title, blocks = _parse_blocks(content, title)
    out = _out(output_path, "document.pdf", ".pdf")
    require("reportlab", "reportlab")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    all_text = title + "\n" + "\n".join(
        b.get("text", "") + " ".join(b.get("items", [])) + " ".join(c for r in b.get("rows", []) for c in r)
        for b in blocks)
    font_name, bold_name, font_file = "Helvetica", "Helvetica-Bold", ""
    need = _needs_unicode_font(all_text)
    has_dev = bool(_DEVANAGARI.search(all_text))
    note = ""
    if need or font_path:
        candidates = _font_candidates(font_path)
        for cand in candidates:
            if not _missing_glyphs(cand, need):
                font_file = cand
                break
        if not font_file and need:
            if not candidates:
                raise ToolError(
                    "text needs a Unicode font (e.g. Devanagari) but none was found. Pass font_path=<.ttf in "
                    "workspace> or set AGENT_PDF_FONT (Windows: C:\\Windows\\Fonts\\Nirmala.ttf).")
            lacking = "".join(sorted(_missing_glyphs(candidates[0], need)))[:10]
            raise ToolError(
                f"no available font has glyphs for: {lacking!r} (tried {len(candidates)} font(s)). "
                "Pass font_path=<.ttf with those glyphs> or set AGENT_PDF_FONT.")
        if not font_file:  # font_path given but text is plain Latin: still use it
            font_file = candidates[0]
    if font_file:
        fname = "AgentFont_" + re.sub(r"\W", "_", Path(font_file).stem)
        try:
            pdfmetrics.registerFont(TTFont(fname, font_file, subfontIndex=0))
        except Exception as e:  # noqa: BLE001
            raise ToolError(f"cannot load font {Path(font_file).name}: {e}")
        pdfmetrics.registerFontFamily(fname, normal=fname, bold=fname, italic=fname, boldItalic=fname)
        font_name = bold_name = fname
    if has_dev:
        note = (" NOTE: reportlab does not shape Devanagari; vowel signs/conjuncts may look imperfect. "
                "For exact rendering use markdown_converter + browser Print to PDF.")

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=font_name, fontSize=11, leading=15)
    heads = {n: ParagraphStyle(f"h{n}", parent=styles["Heading1"], fontName=bold_name, fontSize=sz, leading=sz + 4)
             for n, sz in ((0, 22), (1, 18), (2, 15), (3, 13))}

    def P(text: str, style) -> Paragraph:
        return Paragraph(_pdf_markup(text), style)

    story: List[Any] = []
    if title:
        story += [P(title, heads[0]), Spacer(1, 10)]
    for b in blocks:
        t = b["type"]
        if t == "heading":
            story += [Spacer(1, 6), P(b["text"], heads[b["level"]])]
        elif t == "paragraph":
            story += [P(b["text"], body), Spacer(1, 6)]
        elif t in ("bullets", "numbered"):
            items = [ListItem(P(i, body)) for i in b["items"]]
            story.append(ListFlowable(items, bulletType="bullet" if t == "bullets" else "1", leftIndent=18))
        else:
            rows = b["rows"]
            n = max(len(r) for r in rows)
            data = [[P(c, body) for c in r + [""] * (n - len(r))] for r in rows]
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                     ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                     ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [tbl, Spacer(1, 8)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=title or out.stem).build(story)
    atomic_write_bytes(out, buf.getvalue())
    used = Path(font_file).name if font_file else "Helvetica"
    return f"PDF saved: {out} (font: {used}){note}"


# ---------------------------------------------------------------------------
# pptx_creator
# ---------------------------------------------------------------------------
def _parse_slides(content: str) -> List[Dict[str, Any]]:
    _check_content(content)
    s = content.strip()
    slides: List[Dict[str, Any]]
    data = _try_json(s)
    if data is not None:
        slides = data.get("slides") if isinstance(data, dict) else data  # type: ignore[assignment]
        if not isinstance(slides, list):
            raise ToolError('JSON needs a list of slides: [{"title":..,"bullets":[..],"notes":..}]')
    else:
        slides, cur = [], None
        for line in s.splitlines():
            t = line.strip()
            if not t:
                continue
            m = re.match(r"^#{1,2}\s+(.*)$", t)
            if m:
                cur = {"title": m.group(1), "bullets": []}
                slides.append(cur)
            else:
                if cur is None:
                    cur = {"title": "", "bullets": []}
                    slides.append(cur)
                cur["bullets"].append(re.sub(r"^[-*•]\s+", "", t))
    if not slides:
        raise ToolError("no slides found")
    if len(slides) > MAX_SLIDES:
        raise ToolError(f"more than {MAX_SLIDES} slides")
    out = []
    for sl in slides:
        if not isinstance(sl, dict):
            raise ToolError("every slide must be an object")
        bullets = sl.get("bullets", [])
        if not isinstance(bullets, list):
            raise ToolError("slide bullets must be a list")
        if len(bullets) > 50:
            raise ToolError("a slide can have at most 50 bullets; split it into several slides")
        out.append({"title": _clean(sl.get("title", "")), "bullets": [_clean(b) for b in bullets],
                    "notes": _clean(sl.get("notes", ""))})
    return out


@tool
@guard
def pptx_creator(content: str, output_path: str = "presentation.pptx") -> str:
    """Create a PowerPoint (.pptx) deck and save it in the workspace. Needs: pip install python-pptx.

    Args:
        content: JSON list of slides [{"title":"..","bullets":["..",".."],"notes":".."}] (or {"slides":[...]}), or text
            where each "# Title" line starts a slide and following lines are its bullets.
        output_path: destination .pptx path inside the workspace (default presentation.pptx).
    """
    slides = _parse_slides(content)
    out = _out(output_path, "presentation.pptx", ".pptx")
    pptx = require("pptx", "python-pptx")
    prs = pptx.Presentation()
    for i, sl in enumerate(slides):
        if i == 0 and not sl["bullets"]:
            s = prs.slides.add_slide(prs.slide_layouts[0])
            s.shapes.title.text = sl["title"]
        elif not sl["bullets"]:
            s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
            s.shapes.title.text = sl["title"]
        else:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = sl["title"]
            tf = s.placeholders[1].text_frame
            tf.clear()
            for j, b in enumerate(sl["bullets"]):
                para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                para.text = b
        if sl["notes"]:
            s.notes_slide.notes_text_frame.text = sl["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    atomic_write_bytes(out, buf.getvalue())
    return f"PPTX saved: {out} ({len(slides)} slides)"


# ---------------------------------------------------------------------------
# xlsx_creator
# ---------------------------------------------------------------------------
def _csv_value(cell: str) -> Any:
    """CSV text -> int/float when the cell is a plain number (so Excel can sum it), else the string."""
    c = cell.strip()
    if re.fullmatch(r"[+-]?\d{1,15}", c) and not re.fullmatch(r"[+-]?0\d+", c):
        return int(c)
    if re.fullmatch(r"[+-]?(\d+\.\d+|\.\d+)", c) and len(c) <= 16:
        return float(c)
    return cell


def _parse_sheets(content: str) -> Dict[str, List[List[Any]]]:
    _check_content(content)
    s = content.strip()
    data = _try_json(s)
    if data is not None:
        if isinstance(data, list):
            data = {"Sheet1": data}
        if "sheets" in data and isinstance(data["sheets"], dict):
            data = data["sheets"]
    else:
        try:
            dialect = csv.Sniffer().sniff(s[:4096], delimiters=",;\t|")
            rows = list(csv.reader(io.StringIO(s), dialect))
        except csv.Error:
            try:
                rows = list(csv.reader(io.StringIO(s)))
            except csv.Error as e:
                raise ToolError(f"malformed CSV: {e}")
        data = {"Sheet1": [[_csv_value(c) for c in r] for r in rows]}
    if not data:
        raise ToolError("no sheets found")
    if len(data) > 50:
        raise ToolError("more than 50 sheets")
    total = 0
    sheets: Dict[str, List[List[Any]]] = {}
    for name, rows in data.items():
        if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
            raise ToolError(f"sheet '{name}' must be a list of rows (lists)")
        total += sum(len(r) for r in rows)
        if total > MAX_TABLE_CELLS * 5:
            raise ToolError("too many cells")
        clean = _clean(re.sub(r"[\[\]:*?/\\]", "_", str(name)))[:31] or "Sheet"
        base, k = clean, 2
        while clean.lower() in {n.lower() for n in sheets}:  # Excel sheet names are case-insensitive
            clean = f"{base[:28]}_{k}"
            k += 1
        sheets[clean] = rows
    return sheets


@tool
@guard
def xlsx_creator(content: str, output_path: str = "workbook.xlsx", allow_formulas: bool = False) -> str:
    """Create an Excel (.xlsx) workbook and save it in the workspace. Needs: pip install openpyxl.
    Text cells starting with = + - @ are stored as plain text unless allow_formulas is true (formula injection guard).

    Args:
        content: JSON {"sheets": {"Name": [["h1","h2"],[1,2]]}}, a bare list of rows, or CSV text. First row is bold.
        output_path: destination .xlsx path inside the workspace (default workbook.xlsx).
        allow_formulas: if true, strings starting with "=" are written as real Excel formulas (default false).
    """
    sheets = _parse_sheets(content)
    out = _out(output_path, "workbook.xlsx", ".xlsx")
    openpyxl = require("openpyxl", "openpyxl")
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n_cells = 0
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row, start=1):
                if isinstance(val, (dict, list)):
                    val = json.dumps(val, ensure_ascii=False)
                if isinstance(val, str):
                    val = _clean(val)
                elif not (val is None or isinstance(val, (bool, int, float))):
                    val = _clean(val)
                elif isinstance(val, float) and (val != val or val in (float("inf"), float("-inf"))):
                    val = None
                cell = ws.cell(row=i, column=j)
                if isinstance(val, str) and val[:1] in ("=", "+", "-", "@"):
                    if val[:1] == "=" and allow_formulas:
                        cell.value = val
                    else:
                        cell.value = val
                        cell.data_type = "s"
                else:
                    cell.value = val
                if i == 1:
                    cell.font = Font(bold=True)
                n_cells += 1
        for j in range(1, min(ws.max_column, 200) + 1):  # readable column widths
            longest = max((len(str(ws.cell(row=r, column=j).value or "")) for r in range(1, min(ws.max_row, 200) + 1)),
                          default=8)
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = min(max(longest + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    atomic_write_bytes(out, buf.getvalue())
    return f"XLSX saved: {out} ({len(sheets)} sheet(s), {n_cells} cells)"


# ---------------------------------------------------------------------------
# markdown_converter
# ---------------------------------------------------------------------------
_HTML_TEMPLATE = (
    '<!DOCTYPE html>\n<html lang="en"><head><meta charset="utf-8">'
    '<meta http-equiv="Content-Security-Policy" content="script-src \'none\'; object-src \'none\'">'
    '<meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title>'
    "<style>body{{font-family:'Nirmala UI','Segoe UI',Arial,sans-serif;max-width:800px;margin:2em auto;"
    "padding:0 1em;line-height:1.6}}table{{border-collapse:collapse}}td,th{{border:1px solid #999;padding:4px 8px}}"
    "code,pre{{background:#f4f4f4}}pre{{padding:8px;overflow:auto}}</style></head><body>\n{body}\n</body></html>\n"
)


def _inline(text: str) -> str:
    t = html.escape(text)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    t = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<em>\1</em>", t)
    t = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', t)
    return t


def _basic_markdown(text: str) -> str:
    """Tiny stdlib Markdown subset (headings, lists, code fences, tables, paragraphs, inline). HTML is escaped."""
    out: List[str] = []
    in_code = False
    code: List[str] = []
    list_kind = ""
    para: List[str] = []

    def close_list():
        nonlocal list_kind
        if list_kind:
            out.append(f"</{list_kind}>")
            list_kind = ""

    def flush():
        if para:
            out.append("<p>" + _inline(" ".join(para)) + "</p>")
            para.clear()

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        i += 1
        if s.strip().startswith("```"):
            if in_code:
                out.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
                code.clear()
            else:
                flush()
                close_list()
            in_code = not in_code
            continue
        if in_code:
            code.append(s)
            continue
        st = s.strip()
        m_h = re.match(r"^(#{1,6})\s+(.*)$", st)
        m_b = re.match(r"^[-*+]\s+(.*)$", st)
        m_n = re.match(r"^\d+[.)]\s+(.*)$", st)
        if not st:
            flush()
            close_list()
        elif m_h:
            flush()
            close_list()
            n = len(m_h.group(1))
            out.append(f"<h{n}>{_inline(m_h.group(2))}</h{n}>")
        elif m_b or m_n:
            flush()
            kind = "ul" if m_b else "ol"
            if list_kind != kind:
                close_list()
                out.append(f"<{kind}>")
                list_kind = kind
            out.append("<li>" + _inline((m_b or m_n).group(1)) + "</li>")  # type: ignore[union-attr]
        elif st.startswith("|") and st.endswith("|"):
            flush()
            close_list()
            rows = [st]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            html_rows = []
            for k, r in enumerate(rows):
                cells = [c.strip() for c in r.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                tag = "th" if k == 0 else "td"
                html_rows.append("<tr>" + "".join(f"<{tag}>{_inline(c)}</{tag}>" for c in cells) + "</tr>")
            out.append("<table>" + "".join(html_rows) + "</table>")
        elif st.startswith(">"):
            flush()
            close_list()
            out.append("<blockquote>" + _inline(st.lstrip("> ")) + "</blockquote>")
        else:
            close_list()
            para.append(st)
    if in_code:
        out.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
    flush()
    close_list()
    return "\n".join(out)


_SAFE_URL = re.compile(r"^(https?:|mailto:|#|/|\./|\.\./|[A-Za-z0-9_][^:]*$)", re.I)


def _sanitize_html(body: str) -> str:
    """Neutralise dangerous URL schemes in href/src attributes (javascript:, data:, vbscript:, obfuscated forms)."""
    def fix(m: "re.Match[str]") -> str:
        attr, quote, value = m.group(1), m.group(2), m.group(3)
        plain = re.sub(r"[\x00-\x20]+", "", html.unescape(value))
        if _SAFE_URL.match(plain):
            return m.group(0)
        return f"{attr}={quote}#blocked{quote}"

    return re.sub(r"(?i)\b(href|src)\s*=\s*([\"'])(.*?)\2", fix, body, flags=re.S)


def _markdown_lib_html(md_text: str) -> str:
    import markdown as md  # type: ignore
    from markdown.extensions import Extension  # type: ignore

    class NoRawHTML(Extension):
        def extendMarkdown(self, m):  # raw HTML in the source becomes escaped text
            m.preprocessors.deregister("html_block")
            m.inlinePatterns.deregister("html")

    return md.markdown(md_text, extensions=["tables", "fenced_code", NoRawHTML()])


@tool
@guard
def markdown_converter(content: str, output_path: str = "document.html", title: str = "") -> str:
    """Convert Markdown text to a standalone HTML file saved in the workspace. Uses the 'markdown' package when
    installed (pip install markdown), otherwise a built-in converter. Raw HTML in the input is escaped, unsafe
    link schemes (javascript:, data:) are blocked and the page carries a no-script Content-Security-Policy,
    so the output is safe to open.

    Args:
        content: Markdown text (UTF-8, Hindi is fine).
        output_path: destination .html path inside the workspace (default document.html).
        title: optional <title> of the page (defaults to the first heading).
    """
    _check_content(content)
    content = _clean(content)
    out = _out(output_path, "document.html", ".html")
    try:
        body = _markdown_lib_html(content)
        engine = "markdown"
    except ImportError:
        body = _basic_markdown(content)
        engine = "builtin"
    body = _sanitize_html(body)
    if not title:
        m = re.search(r"^#{1,6}\s+(.+)$", content, re.M)
        title = m.group(1).strip() if m else out.stem
    atomic_write_text(out, _HTML_TEMPLATE.format(title=html.escape(_clean(title)), body=body))
    return f"HTML saved: {out} (engine: {engine})"


DOC_GENERATION_TOOLS = [docx_creator, pdf_creator, pptx_creator, xlsx_creator, markdown_converter]
